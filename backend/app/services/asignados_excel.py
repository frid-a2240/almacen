"""Reporte de "quién tiene asignada" una herramienta: a diferencia del
inventario por empleado (que edita la plantilla real .xlsm), este es un
reporte nuevo sin plantilla en AppSheet — un .xlsx simple generado con
openpyxl directo es suficiente.
"""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font

_ENCABEZADOS = [
    "Número Empleado", "Nombre", "Puesto", "Departamento",
    "Cantidad", "Fecha de Salida", "Número de Vale", "Observaciones",
]
_ANCHOS = [16, 30, 24, 24, 10, 14, 16, 34]
_FILA_ENCABEZADO = 3


def generar_asignados_excel(descripcion_producto: str, codigo_sai_sku: str, filas: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Asignados"

    ws["A1"] = "Herramienta"
    ws["A1"].font = Font(bold=True)
    ws["B1"] = f"{descripcion_producto} ({codigo_sai_sku})"

    for col, texto in enumerate(_ENCABEZADOS, start=1):
        celda = ws.cell(row=_FILA_ENCABEZADO, column=col, value=texto)
        celda.font = Font(bold=True)

    for offset, f in enumerate(filas):
        fila = _FILA_ENCABEZADO + 1 + offset
        ws.cell(row=fila, column=1, value=f["empleado_id"])
        ws.cell(row=fila, column=2, value=f["nombre_de_empleado"])
        ws.cell(row=fila, column=3, value=f["puesto_posicion"])
        ws.cell(row=fila, column=4, value=f["departamento"])
        ws.cell(row=fila, column=5, value=float(f["cantidad"]))
        ws.cell(row=fila, column=6, value=f["fecha"].strftime("%d/%m/%Y") if f["fecha"] else None)
        ws.cell(row=fila, column=7, value=f["numero_de_vale"])
        ws.cell(row=fila, column=8, value=f["observaciones"])

    for letra, ancho in zip("ABCDEFGH", _ANCHOS):
        ws.column_dimensions[letra].width = ancho

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.read()
