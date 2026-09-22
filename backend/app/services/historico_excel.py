"""Registro histórico (Excel simple, sin plantilla) de TODOS los movimientos
de un empleado o de una herramienta a lo largo del tiempo — a diferencia de
inventario_excel.py/asignados_excel.py (que solo muestran lo que se tiene
ACTUALMENTE), aquí se listan también las salidas ya devueltas o traspasadas,
marcadas con su estado, para tener la línea de tiempo completa.
"""
from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font

_FILA_ENCABEZADO = 4


def _fecha_str(f):
    return f.strftime("%d/%m/%Y") if isinstance(f, date) else (f or "")


def _antiguedad(fecha_de_alta):
    """"X años Y meses" desde fecha_de_alta hasta hoy — None si no hay fecha."""
    if not isinstance(fecha_de_alta, date):
        return None
    hoy = date.today()
    meses = (hoy.year - fecha_de_alta.year) * 12 + (hoy.month - fecha_de_alta.month)
    if hoy.day < fecha_de_alta.day:
        meses -= 1
    meses = max(meses, 0)
    anios, meses_resto = divmod(meses, 12)
    partes = []
    if anios:
        partes.append(f"{anios} año{'s' if anios != 1 else ''}")
    if meses_resto or not partes:
        partes.append(f"{meses_resto} mes{'es' if meses_resto != 1 else ''}")
    return " ".join(partes)


def _estado(tipo_movimiento: str, activo: bool) -> str:
    if tipo_movimiento == "SALIDA":
        return "Activo (en resguardo)" if activo else "Ya devuelta / traspasada"
    return "Devolución / traspaso saliente"


def generar_historico_empleado_excel(nombre_empleado: str, id_numero_empleado: str, puesto: str, filas: list[dict]) -> bytes:
    """filas: cada movimiento (SALIDA y ENTRADA) de este empleado, con
    "activo" ya resuelto (saldo actual > 0 para ese SKU)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Registro Histórico"

    ws["A1"] = "Empleado"
    ws["A1"].font = Font(bold=True)
    ws["B1"] = f"{nombre_empleado} ({id_numero_empleado})"
    ws["A2"] = "Puesto"
    ws["A2"].font = Font(bold=True)
    ws["B2"] = puesto or ""

    encabezados = [
        "Fecha", "Tipo", "Folio", "Código SAI SKU", "Descripción",
        "Cantidad", "Número Económico", "Costo Unitario", "Estado", "Observaciones",
    ]
    anchos = [12, 10, 10, 18, 34, 10, 16, 14, 24, 34]
    for col, texto in enumerate(encabezados, start=1):
        celda = ws.cell(row=_FILA_ENCABEZADO, column=col, value=texto)
        celda.font = Font(bold=True)

    for offset, f in enumerate(filas):
        fila = _FILA_ENCABEZADO + 1 + offset
        ws.cell(row=fila, column=1, value=_fecha_str(f.get("fecha")))
        ws.cell(row=fila, column=2, value=f.get("tipo_movimiento"))
        ws.cell(row=fila, column=3, value=f.get("numero_de_vale"))
        ws.cell(row=fila, column=4, value=f.get("sku"))
        ws.cell(row=fila, column=5, value=f.get("descripcion"))
        ws.cell(row=fila, column=6, value=float(f["cantidad"]) if f.get("cantidad") is not None else None)
        ws.cell(row=fila, column=7, value=f.get("numero_economico"))
        ws.cell(row=fila, column=8, value=float(f["costo_unitario"]) if f.get("costo_unitario") else None)
        ws.cell(row=fila, column=9, value=_estado(f.get("tipo_movimiento"), f.get("activo", False)))
        ws.cell(row=fila, column=10, value=f.get("observaciones"))

    for letra, ancho in zip("ABCDEFGHIJ", anchos):
        ws.column_dimensions[letra].width = ancho

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.read()


def generar_historico_producto_excel(descripcion: str, codigo_sai_sku: str, fecha_de_alta, filas: list[dict]) -> bytes:
    """filas: cada movimiento (SALIDA y ENTRADA) de esta herramienta, de
    cualquier empleado que alguna vez la haya tenido, con "activo" resuelto."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Registro Histórico"

    ws["A1"] = "Herramienta"
    ws["A1"].font = Font(bold=True)
    ws["B1"] = f"{descripcion} ({codigo_sai_sku})"
    ws["A2"] = "Fecha de alta"
    ws["A2"].font = Font(bold=True)
    ws["B2"] = _fecha_str(fecha_de_alta)
    ws["A3"] = "Antigüedad"
    ws["A3"].font = Font(bold=True)
    ws["B3"] = _antiguedad(fecha_de_alta) or ""

    encabezados = [
        "Fecha", "Tipo", "Folio", "Empleado", "ID Empleado",
        "Cantidad", "Estado", "Observaciones",
    ]
    anchos = [12, 10, 10, 30, 14, 10, 24, 34]
    for col, texto in enumerate(encabezados, start=1):
        celda = ws.cell(row=_FILA_ENCABEZADO, column=col, value=texto)
        celda.font = Font(bold=True)

    for offset, f in enumerate(filas):
        fila = _FILA_ENCABEZADO + 1 + offset
        ws.cell(row=fila, column=1, value=_fecha_str(f.get("fecha")))
        ws.cell(row=fila, column=2, value=f.get("tipo_movimiento"))
        ws.cell(row=fila, column=3, value=f.get("numero_de_vale"))
        ws.cell(row=fila, column=4, value=f.get("nombre_de_empleado"))
        ws.cell(row=fila, column=5, value=f.get("empleado_id"))
        ws.cell(row=fila, column=6, value=float(f["cantidad"]) if f.get("cantidad") is not None else None)
        ws.cell(row=fila, column=7, value=_estado(f.get("tipo_movimiento"), f.get("activo", False)))
        ws.cell(row=fila, column=8, value=f.get("observaciones"))

    for letra, ancho in zip("ABCDEFGH", anchos):
        ws.column_dimensions[letra].width = ancho

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.read()
